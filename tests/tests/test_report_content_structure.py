import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'app'))

import report.engine as engine_module
from report.engine import ReportEngine, _parse_section_items, _render_section_items, normalize_markdown_text
from report.pdf_export import _clean_text
from report.template_parser import TemplateParser
from report.tables import build_task_table, resolve_report_tasks


def test_template_context_and_multiline_numbered_replacement(tmp_path):
    template_path = tmp_path / 'template.docx'
    output_path = tmp_path / 'output.docx'
    document = Document()
    document.add_paragraph('四、重点产品专项分析')
    document.add_paragraph('4.2 成本异常点排查')
    document.add_paragraph('{{成本异常排查分析}}')
    document.save(template_path)

    parser = TemplateParser(template_path)
    assert parser.placeholder_contexts()['成本异常排查分析'][0]['heading'] == '4.2 成本异常点排查'
    parser.replace_placeholders({
        '成本异常排查分析': '1. 异常判断：单位成本出现波动。\n1.1 数据证据：本月为7.26元/盒。\n1.2 核查范围：需进一步核查工时和费用分摊表。'
    })
    parser.save(output_path)

    rendered = Document(output_path)
    body = [paragraph.text for paragraph in rendered.paragraphs if paragraph.text]
    assert '1. 异常判断：单位成本出现波动。' in body
    assert '1.1 数据证据：本月为7.26元/盒。' in body
    assert all(not text.startswith('• ') for text in body)
    numbered = [paragraph for paragraph in rendered.paragraphs if paragraph.text.startswith(('1. ', '1.1 ', '1.2 '))]
    assert numbered
    assert all(run.bold is False for paragraph in numbered for run in paragraph.runs if run.text)


def test_structured_section_output_is_validated_and_numbered():
    valid = _parse_section_items(
        '{"items":[{"level":1,"title":"异常判断","content":"成本环比上升3.20%。"},'
        '{"level":2,"title":"核查范围","content":"需进一步核查费用分摊表。"}]}'
    )
    assert _render_section_items(valid).splitlines() == [
        '1. 异常判断：成本环比上升3.20%。',
        '1.1 核查范围：需进一步核查费用分摊表。',
    ]
    assert _parse_section_items('1. 不是JSON') == []
    assert _parse_section_items('{"items":[{"level":1,"title":"完整报告","content":"不应使用"}]}') == []


def test_section_generator_uses_json_then_data_fallback(monkeypatch):
    engine = ReportEngine()
    monkeypatch.setattr(engine_module, '_safe_llm_call', lambda *args, **kwargs: '{"items":[{"level":1,"title":"数据结论","content":"本月成本为7.26元/盒。"}]}')
    text = engine._generate_template_section(
        '成本异常排查分析', {'heading': '4.2 成本异常点排查'},
        '板蓝根颗粒', '2026-06', '趋势数据', '',
        [{'level': 1, 'title': '兜底', 'content': '不应出现'}],
    )
    assert text == '1. 数据结论：本月成本为7.26元/盒。'

    monkeypatch.setattr(engine_module, '_safe_llm_call', lambda *args, **kwargs: '模型格式错误')
    fallback = engine._generate_template_section(
        '成本异常排查分析', {'heading': '4.2 成本异常点排查'},
        '板蓝根颗粒', '2026-06', '趋势数据', '',
        [{'level': 1, 'title': '数据兜底', 'content': '需进一步核查原始记录。'}],
    )
    assert fallback == '1. 数据兜底：需进一步核查原始记录。'


def test_safe_llm_call_formats_prompt_when_client_is_available(monkeypatch):
    class FakeClient:
        def chat(self, prompt, system):
            assert '产品：板蓝根颗粒' in prompt
            assert system
            return 'ok'

    monkeypatch.setattr(engine_module, '_HAS_LLM', True)
    monkeypatch.setattr(engine_module, 'llm_client', FakeClient())
    assert engine_module._safe_llm_call('产品：{product}\n{context}', '数据', product='板蓝根颗粒') == 'ok'


def test_numbered_text_survives_report_and_pdf_cleaning():
    value = '1. 总结：成本稳定。\n1.1 证据：环比-2.81%。\n1.1.1 核查：核对工时。'
    assert normalize_markdown_text(value) == value
    assert _clean_text(value) == value


def test_preview_supports_nested_report_numbering():
    source = (Path(__file__).resolve().parents[2] / 'app/static/js/report.js').read_text(encoding='utf-8')
    assert 'const label = level === 1' in source
    assert 'report-numbered-item level-${level}' in source
    assert r'\d+(?:\.\d+){0,2}' in source
    css = (Path(__file__).resolve().parents[2] / 'app/static/css/style.css').read_text(encoding='utf-8')
    assert '.report-numbered-item.level-1' in css
    assert 'font-weight: 400;' in css


def test_report_task_table_and_preview_candidates_share_resolved_tasks():
    llm_tasks = '''[
        {"task_title": "核查金银花采购入库价与合同单价差异", "assignee": {"name": "张伟", "department": "采购部", "role": "采购经理"}, "priority": "high", "deadline": "2026-07-20", "expected_result": "输出采购入库价与合同单价差异核查表，并提交整改措施。", "source": {"finding": "材料成本高于预算"}}
    ]'''
    items = resolve_report_tasks('银黄口服液', '2026-06', llm_tasks)
    rows = build_task_table('银黄口服液', '2026-06', task_items=items)
    assert len(items) == 1
    assert rows[0][1] == items[0]['task_title']
    assert rows[0][2] == items[0]['assignee']['name']
    assert rows[0][5] == items[0]['deadline']
    assert items[0]['source']['finding'] == '材料成本高于预算'


def test_report_task_entry_uses_preview_candidates_and_task_dialog():
    source = (Path(__file__).resolve().parents[2] / 'app/static/js/report.js').read_text(encoding='utf-8')
    assert 'id="generate-report-tasks"' in source
    assert "analysis_scenario: 'report_generated_tasks'" in source
    assert 'prebuilt_tasks: candidates' in source
    assert 'TaskDraftDialog.open(data.tasks || []' in source
    assert '报告清单${submitted}项，已生成${data.tasks_generated}项草稿' in source


def test_report_parameters_refresh_from_shared_selector_state():
    source = (Path(__file__).resolve().parents[2] / 'app/static/js/report.js').read_text(encoding='utf-8')
    assert 'if (!AppState.selectorsReady) return;' in source
    assert 'this.syncSelectors();' in source
    assert "document.getElementById('reportProduct')" in source
    assert "document.getElementById('reportMonth')" in source
    assert 'AppState.currentProduct = product.value;' in source
    assert 'AppState.currentMonth = month.value;' in source
